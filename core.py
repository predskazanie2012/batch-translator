"""
Batch Translator — Translation core.
Supports: Google Cloud Translation API and DeepL.
"""

import os
import time
from pathlib import Path
from typing import Optional, Callable

LANGUAGES = [
    ("Африкаанс",           "af"),
    ("Амхарский",           "am-ET"),
    ("Английский",          "en"),
    ("Арабский",            "ar"),
    ("Армянский",           "hy-AM"),
    ("Бенгальский",         "bn"),
    ("Болгарский",          "bg"),
    ("Венгерский",          "hu"),
    ("Вьетнамский",         "vi"),
    ("Греческий",           "el"),
    ("Грузинский",          "ka-GE"),
    ("Гуджарати",           "gu"),
    ("Датский",             "da"),
    ("Иврит",               "he"),
    ("Индонезийский",       "id"),
    ("Испанский",           "es"),
    ("Итальянский",         "it"),
    ("Йоруба",              "yo-NG"),
    ("Казахский",           "kk-KZ"),
    ("Кантонский",          "yue"),
    ("Китайский (мандарин)", "zh"),
    ("Корейский",           "ko"),
    ("Литовский",           "lt"),
    ("Малайский",           "ms"),
    ("Маратхи",             "mr"),
    ("Немецкий",            "de"),
    ("Нидерландский",       "nl"),
    ("Норвежский",          "no"),
    ("Пенджаби",            "pa"),
    ("Персидский",          "fa-IR"),
    ("Польский",            "pl"),
    ("Португальский",       "pt"),
    ("Румынский",           "ro"),
    ("Русский",             "ru"),
    ("Сербский",            "sr"),
    ("Словацкий",           "sk"),
    ("Словенский",          "sl"),
    ("Суахили",             "sw-KE"),
    ("Тайский",             "th"),
    ("Тамильский",          "ta"),
    ("Телугу",              "te"),
    ("Турецкий",            "tr"),
    ("Узбекский",           "uz-UZ"),
    ("Украинский",          "uk"),
    ("Филиппинский",        "fil"),
    ("Финский",             "fi"),
    ("Французский",         "fr"),
    ("Хинди",               "hi"),
    ("Хорватский",          "hr"),
    ("Чешский",             "cs"),
    ("Шведский",            "sv"),
    ("Японский",            "ja"),
]

# DeepL target language codes (API). Нет в списке — в auto используется Google Cloud.
DEEPL_MAP = {
    "en": "EN-US", "de": "DE", "fr": "FR", "es": "ES", "it": "IT",
    "pt": "PT-BR", "nl": "NL", "pl": "PL", "ru": "RU", "uk": "UK",
    "cs": "CS", "sk": "SK", "ro": "RO", "hu": "HU", "bg": "BG",
    "el": "EL", "hr": "HR", "sl": "SL", "et": "ET", "lv": "LV",
    "lt": "LT", "sv": "SV", "da": "DA", "fi": "FI", "ja": "JA",
    "ko": "KO", "zh-CN": "ZH", "zh": "ZH", "id": "ID", "tr": "TR", "ar": "AR",
    "no": "NB", "nb": "NB",
    "vi": "VI", "th": "TH", "hi": "HI", "bn": "BN", "pa": "PA",
    "te": "TE", "ms": "MS", "ta": "TA", "mr": "MR", "fa": "FA",
    "fa-IR": "FA",
    "uz": "UZ", "sw": "SW", "sr": "SR", "he": "HE", "kk": "KK",
    "uz-UZ": "UZ", "sw-KE": "SW", "kk-KZ": "KK",
    "af": "AF", "am": "AM",
    "am-ET": "AM",
}


class TranslationError(Exception):
    pass


# Соответствие внутренних кодов (файлы, TTS, ISO) ↔ кодам из
# deep_translator.constants.GOOGLE_LANGUAGES_TO_CODES (бесплатный Google).
# Всё, что не перечислено, передаётся как есть — совпадает со списком ошибки
# «No support for the provided language».
_GOOGLE_FREE_LANG_CODE: dict[str, str] = {
    # Региональные коды в интерфейсе, базовые коды в deep_translator.
    "am-ET": "am",
    "fa-IR": "fa",
    "hy-AM": "hy",
    "ka-GE": "ka",
    "kk-KZ": "kk",
    "sw-KE": "sw",
    "uz-UZ": "uz",
    "yo-NG": "yo",
    # Нет кода fil — только filipino → tl
    "fil": "tl",
    # Нет he — только hebrew → iw (устаревший код Google)
    "he": "iw",
    # Нет двухбуквенного zh — только chinese (simplified) → zh-CN
    "zh": "zh-CN",
    # Кантонского (yue) в списке нет; ближайший поддерживаемый — traditional
    "yue": "zh-TW",
    # Javanese: ISO 639-1 часто jv, у Google — jw
    "jv": "jw",
    # Bokmål (nb) в списке Google нет, есть norwegian → no
    "nb": "no",
    # Устаревший ISO для индонезийского
    "in": "id",
}


_GOOGLE_CLOUD_LANG_CODE: dict[str, str] = {
    "am-ET": "am",
    "fa-IR": "fa",
    "hy-AM": "hy",
    "ka-GE": "ka",
    "kk-KZ": "kk",
    "sw-KE": "sw",
    "uz-UZ": "uz",
    "yo-NG": "yo",
    "yue": "zh-TW",
    "zh": "zh-CN",
}


def _map_google_free_lang(code: Optional[str]) -> Optional[str]:
    if not code:
        return code
    return _GOOGLE_FREE_LANG_CODE.get(code, code)


def _map_google_cloud_lang(code: Optional[str]) -> Optional[str]:
    if not code:
        return code
    return _GOOGLE_CLOUD_LANG_CODE.get(code, code)


# ── Low-level translators ────────────────────────────────────────────────

def _split_chunks(text: str, max_len: int = 4000) -> list[str]:
    """Split text into chunks on word boundaries."""
    words = text.split()
    chunks, cur, cur_len = [], [], 0
    for w in words:
        if cur_len + len(w) + 1 > max_len and cur:
            chunks.append(" ".join(cur))
            cur, cur_len = [w], len(w)
        else:
            cur.append(w)
            cur_len += len(w) + 1
    if cur:
        chunks.append(" ".join(cur))
    return chunks


def _gc_translate(client, text: str, target: str, source: Optional[str]) -> str:
    if len(text) > 4500:
        parts = [_gc_translate(client, c, target, source) for c in _split_chunks(text)]
        return " ".join(parts)
    target = _map_google_cloud_lang(target)
    source = _map_google_cloud_lang(source)
    kwargs = dict(values=text, target_language=target, format_="text")
    if source:
        kwargs["source_language"] = source
    return client.translate(**kwargs)["translatedText"]


_UNICODE_REPLACEMENTS = str.maketrans({
    "\u2026": "...",   # … → ...
    "\u2018": "'",    # ' → '
    "\u2019": "'",    # ' → '
    "\u201c": '"',    # " → "
    "\u201d": '"',    # " → "
    "\u2013": "-",    # – → -
    "\u2014": "--",   # — → --
    "\u00a0": " ",    # неразрывный пробел → пробел
})


def _sanitize_for_free(text: str) -> str:
    """Replace chars that break deep_translator's HTTP layer (latin-1 header issue)."""
    return text.translate(_UNICODE_REPLACEMENTS)


def _free_translate(text: str, target: str, source: Optional[str]) -> str:
    raise TranslationError("The unofficial free translator is not supported. Select Google Cloud or DeepL and configure your own credentials.")


def _deepl_translate(translator, text: str, target: str, source: Optional[str]) -> str:
    import traceback as _tb
    import pathlib as _pl
    deepl_target = DEEPL_MAP.get(target)
    if not deepl_target:
        raise TranslationError(f"DeepL не поддерживает язык: {target}")
    deepl_source = DEEPL_MAP.get(source) if source else None
    try:
        result = translator.translate_text(text, target_lang=deepl_target,
                                           source_lang=deepl_source)
        return result.text
    except Exception as e:
        log_path = _pl.Path(__file__).parent / "deepl_error.log"
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(f"\n=== ERROR ===\ntext repr: {repr(text[:120])}\n")
            f.write(_tb.format_exc())
        raise


# ── Document helpers ─────────────────────────────────────────────────────

def _safe_style(doc, name: str) -> str:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return "Normal"


def _safe_table_style(doc, name: str) -> str:
    try:
        doc.styles[name]
        return name
    except Exception:
        return "Table Grid"


# ── BatchTranslator ──────────────────────────────────────────────────────

class BatchTranslator:
    def __init__(self,
                 engine: str = "google_cloud",
                 gc_credentials: Optional[str] = None,
                 deepl_key: Optional[str] = None):
        self.engine = engine
        self._gc_creds = gc_credentials
        self._deepl_key = deepl_key
        self._gc_client = None
        self._deepl_client = None
        self._stop_flag = False

    def init(self):
        """Initialise the chosen engine (raises on auth error)."""
        if self.engine == "google_cloud":
            if self._gc_creds:
                os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = self._gc_creds
            from google.cloud import translate_v2 as translate
            self._gc_client = translate.Client()
        elif self.engine == "deepl":
            import deepl
            self._deepl_client = deepl.Translator(self._deepl_key)
        elif self.engine == "google_free":
            raise TranslationError("Use Google Cloud or DeepL as the translation provider.")

    def translate(self, text: str, target: str, source: Optional[str] = None) -> str:
        if not text or not text.strip():
            return text
        text = _sanitize_for_free(text)
        if self.engine == "google_cloud":
            return _gc_translate(self._gc_client, text, target, source)
        elif self.engine == "google_free":
            t = _map_google_free_lang(target)
            s = _map_google_free_lang(source) if source else None
            return _free_translate(text, t, s)
        elif self.engine == "deepl":
            return _deepl_translate(self._deepl_client, text, target, source)
        raise TranslationError(f"Unknown engine: {self.engine}")

    def translate_file(self,
                       input_path: str,
                       lang_code: str,
                       fmt: str,
                       out_folder: str,
                       src_lang: Optional[str] = None,
                       progress_cb: Optional[Callable[[int, int], None]] = None) -> str:
        inp = Path(input_path)
        out_dir = Path(out_folder) / lang_code
        out_dir.mkdir(parents=True, exist_ok=True)

        if fmt == "docx":
            out_path = out_dir / f"{inp.stem}_{lang_code}.docx"
            self._to_docx(inp, out_path, lang_code, src_lang, progress_cb)
        else:
            out_path = out_dir / f"{inp.stem}_{lang_code}.txt"
            self._to_txt(inp, out_path, lang_code, src_lang, progress_cb)

        return str(out_path)

    # ── Internal converters ──────────────────────────────────────────────

    def _to_docx(self, inp: Path, out_path: Path, target: str,
                 source: Optional[str], progress_cb):
        from docx import Document

        if inp.suffix.lower() == ".txt":
            lines = inp.read_text(encoding="utf-8").splitlines()
            doc_out = Document()
            non_empty = [l for l in lines if l.strip()]
            total, done = len(non_empty), 0
            for line in lines:
                if line.strip():
                    doc_out.add_paragraph(self._tr_sleep(line, target, source))
                    done += 1
                    if progress_cb:
                        progress_cb(done, total)
                else:
                    doc_out.add_paragraph("")
            doc_out.save(str(out_path))
        else:
            doc_in = Document(str(inp))
            doc_out = Document()
            paras = list(doc_in.paragraphs)
            non_empty = [p for p in paras if p.text.strip()]
            total, done = len(non_empty), 0

            for para in paras:
                new_para = doc_out.add_paragraph(style=_safe_style(doc_out, para.style.name))
                if para.text.strip():
                    translated = self._tr_sleep(para.text, target, source)
                    run = new_para.add_run(translated)
                    if para.runs:
                        src_run = para.runs[0]
                        run.bold = src_run.bold
                        run.italic = src_run.italic
                        run.underline = src_run.underline
                        if src_run.font.size:
                            run.font.size = src_run.font.size
                    done += 1
                    if progress_cb:
                        progress_cb(done, total)

            for table in doc_in.tables:
                tbl = doc_out.add_table(rows=len(table.rows), cols=len(table.columns))
                tbl.style = _safe_table_style(doc_out, table.style.name if table.style else "Table Grid")
                for r_i, row in enumerate(table.rows):
                    for c_i, cell in enumerate(row.cells):
                        txt = cell.text.strip()
                        tbl.cell(r_i, c_i).text = self._tr_sleep(txt, target, source) if txt else ""

            doc_out.save(str(out_path))

    def _to_txt(self, inp: Path, out_path: Path, target: str,
                source: Optional[str], progress_cb):
        if inp.suffix.lower() in (".docx", ".doc"):
            from docx import Document
            lines = [p.text for p in Document(str(inp)).paragraphs]
        else:
            lines = inp.read_text(encoding="utf-8").splitlines()

        total = sum(1 for l in lines if l.strip())
        done, result = 0, []
        for line in lines:
            if line.strip():
                result.append(self._tr_sleep(line, target, source))
                done += 1
                if progress_cb:
                    progress_cb(done, total)
            else:
                result.append("")

        out_path.write_text("\n".join(result), encoding="utf-8")

    def _tr_sleep(self, text: str, target: str, source: Optional[str]) -> str:
        result = self.translate(text, target, source)
        time.sleep(0.05)
        return result
